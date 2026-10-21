"""Generate the Excel import template used by the admin panel's bulk upload."""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

OUT = "templates/PCH_Import_Template.xlsx"

HEADER_FILL = PatternFill("solid", fgColor="5B3E8C")
REQ_FILL = PatternFill("solid", fgColor="C9A227")
HEADER_FONT = Font(bold=True, color="FFFFFF")

# (column, required, note)
PRODUCTS = [
    ("sku", True, "Unique code, e.g. PCH-BR-AME-08"),
    ("name_en", True, "Product name in English"),
    ("name_hi", False, "Product name in Hindi"),
    ("category", True, "Pick from list"),
    ("crystals", True, "Comma separated, e.g. Amethyst, Clear Quartz"),
    ("description_en", True, "Full description in English"),
    ("description_hi", False, "Full description in Hindi"),
    ("healing_properties", False, "Comma separated, e.g. Calm, Sleep, Intuition"),
    ("chakras", False, "Comma separated: Root, Sacral, Solar Plexus, Heart, Throat, Third Eye, Crown"),
    ("zodiac", False, "Comma separated, e.g. Pisces, Virgo"),
    ("size", False, "e.g. 8 mm beads / 4 inch / 120 g"),
    ("weight_grams", False, "Number, used for shipping"),
    ("mrp_inr", True, "Number, price before discount"),
    ("sale_price_inr", False, "Number, leave blank if no sale"),
    ("gst_percent", True, "Number, e.g. 3, 5, 12, 18"),
    ("hsn_code", False, "For GST invoice"),
    ("stock_qty", True, "Number"),
    ("is_featured", False, "YES = show in PCH crystal showcase"),
    ("status", True, "Pick from list"),
    ("image_files", False, "Comma separated file names uploaded to the media library, e.g. ame-1.jpg, ame-2.jpg"),
    ("video_file", False, "File name or YouTube/Instagram link"),
]
PRODUCT_SAMPLE = [
    "PCH-BR-AME-08", "Amethyst Bracelet", "एमेथिस्ट ब्रेसलेट", "Bracelet", "Amethyst",
    "Natural 8 mm amethyst beads on elastic cord.", "", "Calm, Sleep, Intuition",
    "Third Eye, Crown", "Pisces, Aquarius", "8 mm beads", 25, 1499, 1199, 3, "7103",
    20, "YES", "Active", "amethyst-bracelet-1.jpg, amethyst-bracelet-2.jpg", "",
]
CATEGORIES = ["Bracelet", "Tumble", "Raw Stone", "Tower / Point", "Pyramid", "Sphere",
              "Cluster / Geode", "Tree", "Mala", "Pendant", "Ring", "Angel / Carving",
              "Kit / Combo", "Other"]

SERVICES = [
    ("code", True, "Unique code, e.g. SRV-HYP-60"),
    ("type", True, "Session / Course / Webinar / Workshop"),
    ("name_en", True, "Name in English"),
    ("name_hi", False, "Name in Hindi"),
    ("short_description_en", True, "1–2 lines for cards"),
    ("description_en", True, "Full description in English"),
    ("description_hi", False, "Full description in Hindi"),
    ("helps_with", False, "Comma separated, e.g. Anxiety, Insomnia, Trauma"),
    ("duration_minutes", True, "Number (for courses: total hours x 60, or leave blank)"),
    ("mode", True, "Online / On-site / Both"),
    ("price_inr", False, "Number. Blank = 'Price on request'"),
    ("trainer", False, "e.g. Priyanka B. Chamoli / Sarika Bhatnagar"),
    ("bookable_online", True, "YES = calendar booking + payment; NO = enquiry form"),
    ("status", True, "Pick from list"),
    ("image_files", False, "Comma separated file names"),
    ("video_file", False, "File name or link"),
]
SERVICE_SAMPLES = [
    ["SRV-HYP", "Session", "Clinical Hypnotherapy", "", "Gentle, guided work with the subconscious mind.", "", "", "Anxiety, Stress, Trauma, Insomnia, Overthinking", 60, "Both", "", "Priyanka B. Chamoli", "YES", "Active", "", ""],
    ["SRV-CHK", "Session", "Chakra Balancing", "", "Realign your seven energy centres.", "", "", "Low energy, Emotional blocks", 60, "Both", "", "Priyanka B. Chamoli", "YES", "Active", "", ""],
    ["SRV-ENA", "Session", "Energy Alignment", "", "", "", "", "", 60, "Both", "", "Priyanka B. Chamoli", "YES", "Active", "", ""],
    ["SRV-EMO", "Session", "Emotional Healing", "", "", "", "", "Grief, Trauma, Relationships", 60, "Both", "", "Priyanka B. Chamoli", "YES", "Active", "", ""],
    ["SRV-PLR", "Session", "Past Life Regression", "", "", "", "", "", 90, "Both", "", "Priyanka B. Chamoli", "YES", "Active", "", ""],
    ["SRV-MOM", "Session", "Prenatal & New-Mother Support", "", "", "", "", "Pregnancy anxiety, Postpartum stress", 60, "Both", "", "Priyanka B. Chamoli", "YES", "Active", "", ""],
    ["CRS-CH-L1", "Course", "Clinical Hypnotherapy Level 1 Certification", "", "Two-day certification incl. chakra balancing & energy alignment.", "", "", "", 1080, "On-site", "", "Sarika Bhatnagar", "NO", "Active", "", ""],
    ["CRS-EKAA", "Course", "EKAA Clinical Hypnotherapy Course", "", "2-day course: history of hypnosis, dream interpretation and more.", "", "", "", "", "On-site", "", "", "NO", "Active", "", ""],
]

MARKETING = [
    ("date", True, "YYYY-MM-DD"),
    ("brand", True, "PCH / Crystal Hub / Both"),
    ("channel", True, "Meta Ads / Google Ads / Instagram influencer / Event / Print / Other"),
    ("campaign", False, "Campaign name"),
    ("amount_inr", True, "Number"),
    ("notes", False, ""),
]


def build_sheet(ws, columns, rows, lists=None):
    for i, (col, req, note) in enumerate(columns, start=1):
        c = ws.cell(row=1, column=i, value=col + (" *" if req else ""))
        c.fill, c.font = (REQ_FILL if req else HEADER_FILL), HEADER_FONT
        n = ws.cell(row=2, column=i, value=note)
        n.font, n.alignment = Font(italic=True, color="777777"), Alignment(wrap_text=True)
        ws.column_dimensions[c.column_letter].width = max(16, min(40, len(note) // 2 + 12))
    for r, row in enumerate(rows, start=3):
        for i, v in enumerate(row, start=1):
            ws.cell(row=r, column=i, value=v)
    ws.freeze_panes = "B3"
    for col_name, options in (lists or {}).items():
        idx = [c[0] for c in columns].index(col_name) + 1
        letter = ws.cell(row=1, column=idx).column_letter
        dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{letter}3:{letter}2000")


wb = Workbook()
readme = wb.active
readme.title = "READ ME"
for line in [
    "Pure Conscious Healing / Pure Crystal Hub — Import Template",
    "",
    "1. Gold headers marked * are required. Purple headers are optional.",
    "2. Row 2 explains each column. Data starts at row 3 — replace the sample rows.",
    "3. Upload images/videos in Admin → Media Library first, then put their file names in image_files.",
    "4. Upload this file in Admin → Products (or Services / Marketing) → Import Excel.",
    "5. Rows with an existing sku/code are updated; new ones are created.",
]:
    readme.append([line])
readme["A1"].font = Font(bold=True, size=14, color="5B3E8C")
readme.column_dimensions["A"].width = 110

build_sheet(wb.create_sheet("Products"), PRODUCTS, [PRODUCT_SAMPLE],
            {"category": CATEGORIES, "status": ["Active", "Draft", "Out of stock"],
             "is_featured": ["YES", "NO"]})
build_sheet(wb.create_sheet("Services"), SERVICES, SERVICE_SAMPLES,
            {"type": ["Session", "Course", "Webinar", "Workshop"],
             "mode": ["Online", "On-site", "Both"], "bookable_online": ["YES", "NO"],
             "status": ["Active", "Draft"]})
build_sheet(wb.create_sheet("Marketing Spend"), MARKETING,
            [["2026-10-01", "Crystal Hub", "Meta Ads", "Diwali crystals", 5000, ""]],
            {"brand": ["PCH", "Crystal Hub", "Both"]})
wb.save(OUT)
print("wrote", OUT)
